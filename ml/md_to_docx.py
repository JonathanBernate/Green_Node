# -*- coding: utf-8 -*-
"""
Conversor sencillo de Markdown a .docx para la guia de GreenNode.
Soporta: encabezados (#..####), parrafos, listas (-, [ ], [x]), tablas,
bloques de codigo (```), citas (>) y negritas basicas **texto**.
"""
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_runs_with_bold(paragraph, text):
    """Divide el texto por **negritas** y agrega runs."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()

    # Estilo base
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]

        # Bloque de codigo
        if line.strip().startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
            # Sombreado gris claro simulado con borde: dejamos solo la fuente mono
            i += 1
            continue

        # Tabla (linea con | ... |)
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[-:\s|]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2  # saltar header y separador
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.text = ""
                add_runs_with_bold(cell.paragraphs[0], h)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = table.add_row().cells
                for j, val in enumerate(row):
                    if j < len(cells):
                        cells[j].text = ""
                        add_runs_with_bold(cells[j].paragraphs[0], val)
            doc.add_paragraph()
            continue

        # Encabezados
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            doc.add_heading(text, level=level)
            i += 1
            continue

        # Separador ---
        if line.strip() == "---":
            i += 1
            continue

        # Cita >
        if line.strip().startswith(">"):
            text = line.strip().lstrip(">").strip()
            p = doc.add_paragraph(style="Intense Quote")
            add_runs_with_bold(p, text)
            i += 1
            continue

        # Lista con checkbox
        m = re.match(r"^\s*-\s*\[( |x|X)\]\s+(.*)$", line)
        if m:
            checked = m.group(1).lower() == "x"
            text = m.group(2).strip()
            p = doc.add_paragraph(style="List Bullet")
            box = p.add_run(("[x] " if checked else "[ ] "))
            box.bold = True
            add_runs_with_bold(p, text)
            i += 1
            continue

        # Lista con vinetas
        m = re.match(r"^(\s*)-\s+(.*)$", line)
        if m:
            text = m.group(2).strip()
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_bold(p, text)
            i += 1
            continue

        # Lista numerada
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            text = m.group(1).strip()
            p = doc.add_paragraph(style="List Number")
            add_runs_with_bold(p, text)
            i += 1
            continue

        # Linea vacia
        if line.strip() == "":
            i += 1
            continue

        # Parrafo normal
        p = doc.add_paragraph()
        add_runs_with_bold(p, line)
        i += 1

    doc.save(docx_path)
    print(f"Guardado: {docx_path}")


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "../docs/GUIA_PENDIENTES.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else "../docs/GUIA_PENDIENTES.docx"
    convert(src, dst)
